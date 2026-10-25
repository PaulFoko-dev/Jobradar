"""DOCX des modèles JobRadar, produits par python-docx avec la même structure que le PDF."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from .common import (
    colon,
    compact_text,
    identity_lines,
    identity_parts,
    language_of,
    letter_date,
    letter_recipient,
    section_views,
)
from .jobradar_html import THEMES, Theme
from .models import RenderError

_PAGE_WIDTH_CM = 21.0
_SIDE_MARGIN_CM = 1.6


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.lstrip("#"))


def set_fonts(rpr_owner: object, font: str) -> None:
    """Police explicite sur les quatre écritures, sans police de thème (qui l'emporterait sinon)."""
    rpr = rpr_owner.get_or_add_rPr()  # type: ignore[attr-defined]
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attribute), None)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), font)


def _bottom_border(style: ParagraphStyle, color: str, size_eighths: int) -> None:
    ppr = style.element.get_or_add_pPr()
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color.lstrip("#"))
    border.append(bottom)
    ppr.append(border)


def _drop_borders(style: ParagraphStyle) -> None:
    ppr = style.element.get_or_add_pPr()
    for existing in ppr.findall(qn("w:pBdr")):
        ppr.remove(existing)


def _configure(document: DocumentObject, theme: Theme, scale: float, *, title: str, author: str, subject: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(_PAGE_WIDTH_CM)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(_SIDE_MARGIN_CM)
    section.right_margin = Cm(_SIDE_MARGIN_CM)

    properties = document.core_properties
    now = datetime.now(UTC).replace(microsecond=0)
    properties.title = title
    properties.author = author
    properties.last_modified_by = author
    properties.subject = subject
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.created = now
    properties.modified = now
    properties.revision = 1

    styles = document.styles
    normal = styles["Normal"]
    set_fonts(normal.element, theme.docx_font)
    normal.font.size = Pt(round(9.8 * scale * 2) / 2)
    normal.font.color.rgb = _rgb(theme.text)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    _drop_borders(title_style)
    set_fonts(title_style.element, theme.docx_font)
    title_style.font.size = Pt(theme.name_size * scale)
    title_style.font.bold = True
    title_style.font.color.rgb = _rgb(theme.text)
    title_style.paragraph_format.space_after = Pt(2)

    heading = styles["Heading 1"]
    set_fonts(heading.element, theme.docx_font)
    heading.font.size = Pt(round((11.5 if theme.id == "moderne-ats" else 10.5) * scale * 2) / 2)
    heading.font.bold = True
    heading.font.all_caps = theme.id != "moderne-ats"
    heading.font.color.rgb = _rgb(theme.accent)
    heading.paragraph_format.space_before = Pt(10 * scale)
    heading.paragraph_format.space_after = Pt(4 * scale)
    heading.paragraph_format.keep_with_next = True
    _bottom_border(heading, theme.rule, 6)


def _paragraph(document: DocumentObject, text: str = "", *, style: str | None = None) -> Paragraph:
    paragraph = document.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    return paragraph


def _run(paragraph: Paragraph, text: str, *, bold: bool = False, italic: bool = False, color: str | None = None,
         size: float | None = None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    if color:
        run.font.color.rgb = _rgb(color)
    if size:
        run.font.size = Pt(size)


def _identity(document: DocumentObject, identity: dict, language: str, headline: str, theme: Theme,
              scale: float) -> None:
    name, contacts, links = identity_lines(identity, language)
    if name:
        _paragraph(document, name, style="Title")
    if headline:
        paragraph = _paragraph(document)
        paragraph.paragraph_format.space_after = Pt(2)
        _run(paragraph, headline, bold=True, color=theme.accent if theme.id != "sobre-ats" else theme.text,
             size=round(11.5 * scale * 2) / 2)
    lines = [line for line in (contacts, links) if line]
    for position, line in enumerate(lines):
        paragraph = _paragraph(document)
        _run(paragraph, " · ".join(line), color=theme.muted, size=round(8.8 * scale * 2) / 2)
        paragraph.paragraph_format.space_after = Pt(8 * scale) if position == len(lines) - 1 else Pt(0)


def _save(document: DocumentObject) -> bytes:
    stream = BytesIO()
    try:
        document.save(stream)
    except Exception as exc:
        raise RenderError("Le DOCX n'a pas pu être produit.") from exc
    return stream.getvalue()


def cv_docx(composition: dict, identity: dict, theme_id: str | Theme, *, title: str, author: str, scale: float = 1.0) -> bytes:
    theme = theme_id if isinstance(theme_id, Theme) else THEMES[theme_id]
    language = language_of(composition)
    document = Document()
    _configure(document, theme, scale, title=title, author=author,
               subject="Curriculum vitae" if language == "en" else "Curriculum vitæ")
    _identity(document, identity, language, compact_text(composition.get("headline")), theme, scale)
    summary = compact_text(composition.get("summary"))
    if summary:
        _paragraph(document, summary)
    text_width = Cm(_PAGE_WIDTH_CM - 2 * _SIDE_MARGIN_CM)
    small = round(9.2 * scale * 2) / 2
    for view in section_views(composition):
        _paragraph(document, view.title, style="Heading 1")
        for position, entry in enumerate(view.entries):
            head = _paragraph(document)
            head.paragraph_format.keep_with_next = True
            gap = 2.5 if view.type == "certification" else 5
            head.paragraph_format.space_before = Pt(gap * scale if position else 0)
            _run(head, entry.title, bold=True)
            inline = view.type == "certification" and bool(entry.subtitle)
            if inline:
                _run(head, f" — {entry.subtitle}", color=theme.muted)
            if entry.period:
                head.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                _run(head, "\t" + entry.period, color=theme.muted, size=small)
            if entry.subtitle and not inline:
                sub = _paragraph(document)
                sub.paragraph_format.keep_with_next = bool(entry.bullets or entry.details or entry.tech)
                _run(sub, entry.subtitle, italic=theme.id != "technique-projets",
                     bold=theme.id == "technique-projets",
                     color=theme.accent if theme.id == "technique-projets" else theme.muted)
            for detail in [*entry.details, *entry.tech]:
                _run(_paragraph(document), detail, color=theme.muted, size=small)
            for bullet in entry.bullets:
                item = _paragraph(document, bullet, style="List Bullet")
                item.paragraph_format.space_before = Pt(1)
        for label, text in view.lines:
            line = _paragraph(document)
            line.paragraph_format.space_before = Pt(2)
            if label:
                _run(line, f"{label}{colon(language)}", bold=True, color=theme.accent if theme.id == "technique-projets" else None)
            _run(line, text)
    return _save(document)


def letter_docx(letter: dict, identity: dict, theme_id: str | Theme, *, title: str, author: str) -> bytes:
    theme = theme_id if isinstance(theme_id, Theme) else THEMES[theme_id]
    language = language_of(letter)
    document = Document()
    _configure(document, theme, 1.0, title=title, author=author,
               subject="Cover letter" if language == "en" else "Lettre de motivation")
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Normal"].paragraph_format.line_spacing = 1.25
    _identity(document, identity, language, "", theme, 1.0)
    document.paragraphs[-1].paragraph_format.space_after = Pt(20)
    name, _ = identity_parts(identity, language)
    recipient = letter_recipient(letter)
    for line in recipient:
        paragraph = _paragraph(document, line)
        paragraph.paragraph_format.left_indent = Cm(9.5)
    when = _paragraph(document, letter_date(language, compact_text(identity.get("city"))))
    when.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    when.paragraph_format.space_before = Pt(12 if recipient else 0)
    when.paragraph_format.space_after = Pt(20)

    def block(text: str, *, after: float = 9, bold: bool = False) -> None:
        paragraph = _paragraph(document)
        paragraph.paragraph_format.space_after = Pt(after)
        _run(paragraph, text, bold=bold)

    subject = compact_text(letter.get("subject"))
    if subject:
        block(f"{'Subject:' if language == 'en' else 'Objet :'} {subject}", after=16, bold=True)
    if compact_text(letter.get("greeting")):
        block(compact_text(letter.get("greeting")), after=10)
    paragraphs = letter.get("paragraphs")
    for text in paragraphs if isinstance(paragraphs, list) else []:
        if compact_text(text):
            block(compact_text(text))
    if compact_text(letter.get("closing")):
        block(compact_text(letter.get("closing")))
    if name:
        signature = _paragraph(document)
        signature.paragraph_format.space_before = Pt(14)
        _run(signature, name, bold=True)
    return _save(document)
