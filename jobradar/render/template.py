"""Analyse et remplissage des modèles DOCX fournis par l'utilisateur (CV et lettre).

Les indices de la cartographie désignent les paragraphes d'une « histoire » (corps du document, ou en-tête
de page) dans l'ordre du document, y compris ceux des cellules de tableau et des zones de texte : un modèle
à deux colonnes en tableau reste ainsi cartographiable et remplissable.
"""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.styles import BabelFish
from docx.text.paragraph import Paragraph
from lxml import etree

from .common import (
    CV_SECTION_TYPES,
    Entry,
    SectionView,
    colon,
    compact_text,
    identify_section,
    identity_lines,
    identity_parts,
    language_of,
    letter_date,
    letter_recipient,
    section_title,
    section_views,
)
from .models import RenderError, TemplateMap

_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_MC_FALLBACK = f"{{{_MC}}}Fallback"
_MC_ALTERNATE = f"{{{_MC}}}AlternateContent"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_TXBX = qn("w:txbxContent")
_CONTAINERS = {qn("w:tc"): "table", _TXBX: "textbox", qn("w:body"): "body", qn("w:hdr"): "header",
               qn("w:ftr"): "footer"}
_DRAWING_TAGS = (qn("w:drawing"), qn("w:pict"), _MC_ALTERNATE)

_CONTACT_RE = re.compile(
    r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+|(?:\+|\b0)\d[\d .()-]{7,}\d|https?://|www\.|linkedin\.com|github\.com",
    re.IGNORECASE,
)
_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b|\b(?:aujourd.hui|present|en cours|à ce jour)\b", re.IGNORECASE)
_BULLET_TEXT_RE = re.compile(r"^\s*([•●▪■◦\-–*·➢✓])\s+")
_PAGE_NUMBER_RE = re.compile(r"^(page\s*)?\d+(\s*(/|sur|of)\s*\d+)?$", re.IGNORECASE)
_LIST_STYLE_RE = re.compile(r"list|liste|puce|bullet", re.IGNORECASE)
_GREETING_RE = re.compile(r"^(madame|monsieur|mesdames|messieurs|cher|chère|dear|bonjour|hello|to whom)\b", re.I)
_SUBJECT_RE = re.compile(r"^\s*(objet|subject|re)\s*:\s*", re.IGNORECASE)
_CLOSING_RE = re.compile(r"agr[ée]er|salutations|cordialement|sincerely|regards|respectueusement|faithfully", re.I)
_DATE_LINE_RE = re.compile(
    r"\ble\s+\d{1,2}(?:er)?\s+\w+\s+\d{4}\b|\b\d{1,2}/\d{1,2}/\d{2,4}\b|"
    r"\b(?:january|february|march|april|may|june|july|august|september|october|november|december)\s+\d{1,2},\s*\d{4}",
    re.IGNORECASE,
)

LETTER_TYPES = ("letter_date", "letter_recipient", "letter_subject", "letter_greeting", "letter_body",
                "letter_closing", "letter_signature")


# --- Lecture de la structure -------------------------------------------------------------------------


def story_root(document: DocumentObject, part: str) -> tuple[etree._Element, object] | None:
    if part == "body":
        return document.element.body, document._body
    section = document.sections[0]
    if part == "header" and not section.header.is_linked_to_previous:
        return section.header._element, section.header
    if part == "first_page_header" and section.different_first_page_header_footer \
            and not section.first_page_header.is_linked_to_previous:
        return section.first_page_header._element, section.first_page_header
    return None


def story_paragraphs(root: etree._Element) -> list[etree._Element]:
    """Paragraphes dans l'ordre du document, sans les copies de secours (mc:Fallback) des zones de texte."""
    return [p for p in root.iter(qn("w:p")) if not any(a.tag == _MC_FALLBACK for a in p.iterancestors())]


def container_of(p: etree._Element) -> etree._Element:
    for ancestor in p.iterancestors():
        if ancestor.tag in _CONTAINERS:
            return ancestor
    return p.getparent()


def location_of(p: etree._Element) -> str:
    return _CONTAINERS.get(container_of(p).tag, "body")


def paragraph_text(p: etree._Element) -> str:
    """Texte propre au paragraphe (sans celui de ses zones de texte imbriquées ni des suppressions)."""
    parts: list[str] = []
    for node in p.iter(qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")):
        if _nested_in(node, p):
            continue
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        else:
            parts.append("\n")
    return "".join(parts)


def _nested_in(node: etree._Element, p: etree._Element) -> bool:
    for ancestor in node.iterancestors():
        if ancestor is p:
            return False
        if ancestor.tag in {_TXBX, qn("w:del")}:
            return True
    return False


def _style_id(p: etree._Element) -> str | None:
    ppr = p.find(qn("w:pPr"))
    style = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return style.get(qn("w:val")) if style is not None else None


def _style_name(document: DocumentObject, p: etree._Element) -> str:
    style_id = _style_id(p)
    if style_id is None:
        return "Normal"
    for style in document.styles.element.iter(qn("w:style")):
        if style.get(qn("w:styleId")) == style_id:
            name = style.find(qn("w:name"))
            return BabelFish.internal2ui(name.get(qn("w:val"))) if name is not None else style_id
    return style_id


def _first_text_run(p: etree._Element) -> etree._Element | None:
    for run in p.iter(qn("w:r")):
        if _nested_in(run, p):
            continue
        if any((t.text or "").strip() for t in run.findall(qn("w:t"))):
            return run
    return None


def _run_signature(p: etree._Element) -> str:
    run = _first_text_run(p)
    rpr = run.find(qn("w:rPr")) if run is not None else None
    if rpr is None:
        return ""
    keep = [deepcopy(child) for child in rpr if child.tag not in {qn("w:lang"), qn("w:noProof")}]
    return "".join(etree.tostring(child, encoding="unicode") for child in keep)


def _signature(p: etree._Element) -> tuple[str | None, str]:
    return _style_id(p), _run_signature(p)


def _paragraph_style(document: DocumentObject, p: etree._Element) -> dict:
    paragraph = Paragraph(p, document._body)
    run = _first_text_run(p)
    font: dict = {"name": None, "size_pt": None, "bold": None, "italic": None, "color": None}
    if run is not None:
        from docx.text.run import Run

        wrapped = Run(run, paragraph)
        font = {
            "name": wrapped.font.name,
            "size_pt": wrapped.font.size.pt if wrapped.font.size is not None else None,
            "bold": wrapped.bold,
            "italic": wrapped.italic,
            "color": str(wrapped.font.color.rgb) if wrapped.font.color.type is not None
            and wrapped.font.color.rgb is not None else None,
        }
    alignment = p.find(f"{qn('w:pPr')}/{qn('w:jc')}")
    return {
        "paragraph_style": _style_name(document, p),
        "alignment": alignment.get(qn("w:val")) if alignment is not None else None,
        "font": font,
    }


def _is_bullet(document: DocumentObject, p: etree._Element, text: str) -> bool:
    if p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None:
        return True
    return bool(_LIST_STYLE_RE.search(_style_name(document, p))) or bool(_BULLET_TEXT_RE.match(text))


def _has_drawing(p: etree._Element) -> bool:
    return any(True for tag in _DRAWING_TAGS for _ in p.iter(tag))


@dataclass(frozen=True)
class _Story:
    part: str
    paragraphs: list[etree._Element]
    texts: list[str]


def _stories(document: DocumentObject) -> list[_Story]:
    stories: list[_Story] = []
    for part in ("body", "first_page_header", "header"):
        root = story_root(document, part)
        if root is None:
            continue
        paragraphs = story_paragraphs(root[0])
        stories.append(_Story(part, paragraphs, [compact_text(paragraph_text(p)) for p in paragraphs]))
    return stories


def _font_size(document: DocumentObject, p: etree._Element) -> float:
    run = _first_text_run(p)
    size = run.find(f"{qn('w:rPr')}/{qn('w:sz')}") if run is not None else None
    if size is not None and (size.get(qn("w:val")) or "").isdigit():
        return int(size.get(qn("w:val"))) / 2
    return 0.0


# --- Cartographie --------------------------------------------------------------------------------------


def _map_sections(document: DocumentObject, story: _Story) -> list[dict]:
    paragraphs, texts = story.paragraphs, story.texts
    recognized = [(index, identify_section(text)) for index, text in enumerate(texts)]
    recognized = [(index, kind) for index, kind in recognized if kind is not None and len(texts[index]) <= 45]
    heading_signatures = {_signature(paragraphs[index]) for index, _ in recognized}
    content_signatures = set()
    for index, _ in recognized:
        following = next((i for i in range(index + 1, len(texts)) if texts[i]), None)
        if following is not None and identify_section(texts[following]) is None:
            content_signatures.add(_signature(paragraphs[following]))
    boundaries: dict[int, str] = dict(recognized)
    for index, text in enumerate(texts):
        if index in boundaries or not text or len(text) > 45 or len(text.split()) > 5 or text.endswith("."):
            continue
        signature = _signature(paragraphs[index])
        if signature in heading_signatures and signature not in content_signatures \
                and not _CONTACT_RE.search(text):
            boundaries[index] = "other"

    sections: list[dict] = []
    ordered = sorted(boundaries)
    for position, start in enumerate(ordered):
        container = container_of(paragraphs[start])
        limit = ordered[position + 1] if position + 1 < len(ordered) else len(paragraphs)
        end = start + 1
        while end < limit and container_of(paragraphs[end]) is container:
            end += 1
        content = next((paragraphs[i] for i in range(start + 1, end) if texts[i]), None)
        sections.append({
            "type": boundaries[start],
            "title": texts[start],
            "start": start,
            "end": end,
            "part": story.part,
            "container": location_of(paragraphs[start]),
            "style": {
                "heading": _paragraph_style(document, paragraphs[start]),
                "content": _paragraph_style(document, content) if content is not None else {},
            },
        })
    return sections


def _covered(sections: list[dict]) -> set[int]:
    covered: set[int] = set()
    for section in sections:
        covered.update(range(int(section["start"]), int(section["end"])))
    return covered


def _identity_refs(document: DocumentObject, stories: list[_Story], sections: list[dict],
                   reserved: set[int]) -> dict:
    """Repère nom, accroche, résumé et coordonnées hors des sections, dans le corps puis l'en-tête."""
    header: dict = {"location": "absent", "name": None, "headline": None, "summary": None, "contacts": []}
    covered = _covered(sections)
    for story in stories:
        free = [index for index, text in enumerate(story.texts)
                if text and (story.part != "body" or (index not in covered and index not in reserved))]
        contacts = [index for index in free if _CONTACT_RE.search(story.texts[index])]
        if contacts and not header["contacts"]:
            header["contacts"] = [{"part": story.part, "index": index} for index in contacts]
            header["location"] = location_of(story.paragraphs[contacts[0]]) if story.part == "body" else "header"
        if header["name"] is None:
            candidates = [index for index in free if index not in contacts and len(story.texts[index].split()) <= 5
                          and not _DATE_LINE_RE.search(story.texts[index])]
            if candidates:
                def weight(index: int, paragraphs: list[etree._Element] = story.paragraphs) -> tuple[int, float, int]:
                    style = _style_name(document, paragraphs[index])
                    titled = int(bool(re.search(r"title|titre", style, re.IGNORECASE)))
                    return titled, _font_size(document, paragraphs[index]), -index

                name = max(candidates, key=weight)
                header["name"] = {"part": story.part, "index": name}
                container = container_of(story.paragraphs[name])
                after = [index for index in free if index > name and index not in contacts
                         and container_of(story.paragraphs[index]) is container]
                if after and len(story.texts[after[0]].split()) <= 12 and after[0] == name + 1:
                    header["headline"] = {"part": story.part, "index": after[0]}
                if story.part == "body":
                    summary = next((index for index in after if len(story.texts[index].split()) >= 12), None)
                    if summary is not None:
                        header["summary"] = {"part": story.part, "index": summary}
    return header


def _map_letter(stories: list[_Story]) -> tuple[list[dict], set[int]]:
    body = next(story for story in stories if story.part == "body")
    texts = body.texts
    greeting = next((i for i, text in enumerate(texts) if _GREETING_RE.match(text) and len(text.split()) <= 8), None)
    if greeting is None:
        return [], set()
    closing = next((i for i in range(greeting + 1, len(texts)) if _CLOSING_RE.search(texts[i])), None)
    if closing is None:
        return [], set()
    sections: list[dict] = []

    def add(kind: str, start: int, end: int | None = None) -> None:
        sections.append({"type": kind, "title": "", "start": start, "end": (end if end is not None else start + 1),
                         "part": "body", "container": location_of(body.paragraphs[start]), "style": {}})

    subject = next((i for i in range(greeting) if _SUBJECT_RE.match(texts[i])), None)
    dates = [i for i in range(greeting) if _DATE_LINE_RE.search(texts[i]) and len(texts[i].split()) <= 8]
    if dates:
        add("letter_date", dates[0])
    if subject is not None:
        add("letter_subject", subject)
    add("letter_greeting", greeting)
    filled = [i for i in range(greeting + 1, closing) if texts[i]]
    if filled:
        add("letter_body", filled[0], filled[-1] + 1)
    add("letter_closing", closing)
    signature = next((i for i in range(closing + 1, len(texts)) if texts[i] and len(texts[i].split()) <= 5), None)
    if signature is not None:
        add("letter_signature", signature)
    reserved = {index for section in sections for index in range(section["start"], section["end"])}
    # Le bloc destinataire : lignes libres entre les coordonnées et l'objet, repérées une fois l'identité connue.
    return sections, reserved


def _informative(text: str) -> bool:
    """Un simple numéro de page en en-tête ou pied de page ne gêne pas un ATS."""
    return bool(text) and not _PAGE_NUMBER_RE.match(text)


def _ats_issues(document: DocumentObject, stories: list[_Story], header: dict, sections: list[dict],
                is_letter: bool) -> list[str]:
    issues: list[str] = []
    body_root = document.element.body
    for section in document.sections:
        columns = section._sectPr.find(qn("w:cols"))
        if columns is not None:
            number = columns.get(qn("w:num"))
            if (number and number.isdigit() and int(number) > 1) or len(columns.findall(qn("w:col"))) > 1:
                issues.append("Le document est mis en page sur plusieurs colonnes : un ATS peut lire les colonnes "
                              "dans le désordre et mélanger les sections.")
                break
    layout_tables = 0
    for table in body_root.iter(qn("w:tbl")):
        for row in table.iter(qn("w:tr")):
            filled = [cell for cell in row.findall(qn("w:tc"))
                      if any(compact_text(paragraph_text(p)) for p in cell.iter(qn("w:p")))]
            if len(filled) >= 2:
                layout_tables += 1
                break
    if layout_tables:
        issues.append(f"Le document utilise {layout_tables} tableau(x) de mise en page (contenus côte à côte) : "
                      "les ATS lisent souvent ces cellules dans le désordre.")
    body = next(story for story in stories if story.part == "body")
    boxes = [box for box in body_root.iter(_TXBX) if not any(a.tag == _MC_FALLBACK for a in box.iterancestors())]
    if boxes:
        detail = ""
        in_box = {i for i, p in enumerate(body.paragraphs) if location_of(p) == "textbox"}
        refs = [ref for ref in [header.get("name"), *header.get("contacts", [])] if ref and ref["part"] == "body"]
        if any(ref["index"] in in_box for ref in refs):
            detail = " ; le nom ou les coordonnées s'y trouvent"
        elif any(section["start"] in in_box for section in sections):
            detail = " ; des sections du CV s'y trouvent"
        issues.append(f"Le document contient {len(boxes)} zone(s) de texte flottante(s){detail} : leur contenu est "
                      "souvent ignoré par les ATS.")
    page_parts = [story for story in stories if story.part != "body"]
    if any(ref and ref["part"] != "body" for ref in [header.get("name"), *header.get("contacts", [])]):
        issues.append("Le nom ou les coordonnées sont dans l'en-tête de page : de nombreux ATS ignorent les en-têtes "
                      "et pieds de page.")
    elif any(_informative(text) for story in page_parts for text in story.texts):
        issues.append("L'en-tête de page contient du texte, ignoré par de nombreux ATS.")
    section = document.sections[0]
    if not section.footer.is_linked_to_previous and any(
            _informative(compact_text(paragraph_text(p))) for p in story_paragraphs(section.footer._element)):
        issues.append("Le pied de page contient du texte, ignoré par de nombreux ATS.")
    text_width = (section.page_width or 0) - (section.left_margin or 0) - (section.right_margin or 0)
    floating = sum(1 for anchor in body_root.iter(f"{_WP}anchor") if any(True for _ in anchor.iter(_A_BLIP)))
    large = 0
    for inline in body_root.iter(f"{_WP}inline"):
        extent = inline.find(f"{_WP}extent")
        if extent is not None and extent.get("cx", "0").isdigit() and text_width \
                and int(extent.get("cx")) > 0.4 * text_width:
            large += 1
    if floating or large:
        issues.append("Le document contient une image flottante ou de grande taille : un texte porté par une image "
                      "n'est pas lu par les ATS (une petite photo d'identité reste acceptable).")
    if not is_letter and not any(section["type"] in CV_SECTION_TYPES for section in sections):
        issues.append("Aucun titre de section standard (Expérience professionnelle, Formation, Compétences…) n'a été "
                      "reconnu : un ATS ne saura pas découper le CV.")
    return list(dict.fromkeys(issues))


def analyze_docx_template(path: Path) -> TemplateMap:
    """Cartographie les sections, l'emplacement des coordonnées et les risques ATS d'un modèle DOCX."""
    source_path = Path(path)
    try:
        if not source_path.is_file():
            raise OSError
        document = Document(str(source_path))
        stories = _stories(document)
    except Exception as exc:
        raise RenderError("Le modèle DOCX ne peut pas être lu.") from exc
    try:
        return _analyze(document, stories, source_path)
    except RenderError:
        raise
    except Exception as exc:
        raise RenderError("Le modèle DOCX n'a pas pu être analysé.") from exc


def _analyze(document: DocumentObject, stories: list[_Story], source_path: Path) -> TemplateMap:
    body = next(story for story in stories if story.part == "body")
    sections = _map_sections(document, body)
    is_letter = False
    reserved: set[int] = set()
    if not any(section["type"] in CV_SECTION_TYPES for section in sections):
        letter_sections, reserved = _map_letter(stories)
        if letter_sections:
            sections, is_letter = letter_sections, True
    header = _identity_refs(document, stories, sections, reserved)
    if is_letter:
        header["headline"] = header["summary"] = None
        known = reserved | {ref["index"] for ref in [header["name"], *header["contacts"]]
                            if ref and ref["part"] == "body"}
        # Destinataire : lignes libres entre les coordonnées et l'objet (ou la formule d'appel), date exclue.
        limit = min(section["start"] for section in sections
                    if section["type"] in {"letter_subject", "letter_greeting"})
        last_identity = max({ref["index"] for ref in [header["name"], *header["contacts"]]
                             if ref and ref["part"] == "body"} & set(range(limit)), default=-1)
        recipient = [i for i in range(last_identity + 1, limit) if body.texts[i] and i not in known]
        if recipient:
            sections.append({"type": "letter_recipient", "title": "", "start": recipient[0], "end": recipient[-1] + 1,
                             "part": "body", "container": location_of(body.paragraphs[recipient[0]]), "style": {}})
            sections.sort(key=lambda section: section["start"])
    issues = _ats_issues(document, stories, header, sections, is_letter)
    return TemplateMap(source_path, sections, header, issues, not issues)


# --- Écriture ------------------------------------------------------------------------------------------


def _clean_ppr(ppr: etree._Element | None) -> etree._Element | None:
    if ppr is None:
        return None
    copy = deepcopy(ppr)
    for tag in ("w:sectPr", "w:pPrChange"):
        for node in copy.findall(qn(tag)):
            copy.remove(node)
    return copy


def _clean_rpr(rpr: etree._Element | None, *, bold: bool | None = None) -> etree._Element | None:
    copy = deepcopy(rpr) if rpr is not None else None
    if copy is not None:
        for node in copy.findall(qn("w:rPrChange")):
            copy.remove(node)
    if bold is not None:
        if copy is None:
            copy = OxmlElement("w:rPr")
        for tag in ("w:b", "w:bCs"):
            for node in copy.findall(qn(tag)):
                copy.remove(node)
        element = OxmlElement("w:b")
        if not bold:
            element.set(qn("w:val"), "0")
        style = copy.find(qn("w:rStyle"))
        if style is not None:
            style.addnext(element)
        else:
            copy.insert(0, element)
    return copy


def _append_run(p: etree._Element, parent: object, text: str, rpr: etree._Element | None) -> None:
    run = Paragraph(p, parent).add_run(text)
    if rpr is not None:
        existing = run._r.find(qn("w:rPr"))
        if existing is not None:
            run._r.remove(existing)
        run._r.insert(0, deepcopy(rpr))


def _clear_text(p: etree._Element) -> None:
    """Retire le texte du paragraphe mais garde ses propriétés et ses dessins ancrés (photo, zone de texte)."""
    for child in list(p):
        if child.tag == qn("w:pPr"):
            continue
        if child.tag == qn("w:r") and any(True for tag in _DRAWING_TAGS for _ in child.iter(tag)):
            for node in list(child):
                if node.tag in {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")}:
                    child.remove(node)
            continue
        p.remove(child)


@dataclass
class _Model:
    """Paragraphe d'exemple : propriétés de paragraphe et mises en forme de ses segments de texte."""

    ppr: etree._Element | None
    rprs: list[etree._Element | None]
    tab_rpr: etree._Element | None
    has_tab: bool
    text_bullet: str = ""

    @property
    def first(self) -> etree._Element | None:
        return self.rprs[0] if self.rprs else None

    def second(self) -> etree._Element | None:
        return self.rprs[1] if len(self.rprs) > 1 else self.first


def _model(p: etree._Element) -> _Model:
    rprs: list[etree._Element | None] = []
    signatures: list[str] = []
    tab_rpr = None
    seen_tab = False
    for run in p.iter(qn("w:r")):
        if _nested_in(run, p):
            continue
        if run.find(qn("w:tab")) is not None:
            seen_tab = True
        if not any((t.text or "").strip() for t in run.findall(qn("w:t"))):
            continue
        rpr = run.find(qn("w:rPr"))
        if seen_tab and tab_rpr is None:
            tab_rpr = rpr
            continue
        signature = etree.tostring(rpr, encoding="unicode") if rpr is not None else ""
        if signature not in signatures and not seen_tab:
            signatures.append(signature)
            rprs.append(rpr)
    match = _BULLET_TEXT_RE.match(paragraph_text(p))
    return _Model(_clean_ppr(p.find(qn("w:pPr"))), rprs or [None], tab_rpr, seen_tab,
                  f"{match.group(1)} " if match else "")


class _Writer:
    """Insère des paragraphes après une ancre, en clonant les modèles d'une section."""

    def __init__(self, parent: object, anchor: etree._Element) -> None:
        self.parent = parent
        self.anchor = anchor
        self.created: list[etree._Element] = []

    def paragraph(self, model: _Model, segments: list[tuple[str, etree._Element | None]]) -> etree._Element:
        p = OxmlElement("w:p")
        if model.ppr is not None:
            p.append(deepcopy(model.ppr))
        for text, rpr in segments:
            if text:
                _append_run(p, self.parent, text, rpr)
        self.anchor.addnext(p)
        self.anchor = p
        self.created.append(p)
        return p

    def clone(self, element: etree._Element) -> None:
        copy = deepcopy(element)
        self.anchor.addnext(copy)
        self.anchor = copy
        self.created.append(copy)


@dataclass
class _SectionModels:
    title: _Model
    detail: _Model | None
    meta: _Model | None
    bullet: _Model | None
    spacer: etree._Element | None
    line: _Model


def _section_models(document: DocumentObject, paragraphs: list[etree._Element], texts: list[str],
                    fallback: _Model) -> _SectionModels:
    filled = [(p, text) for p, text in zip(paragraphs, texts, strict=True) if text]
    bullets = [p for p, text in filled if _is_bullet(document, p, text)]
    plain = [(p, text) for p, text in filled if p not in bullets]
    title = _model(plain[0][0]) if plain else (_model(filled[0][0]) if filled else fallback)
    detail = meta = None
    for p, text in plain[1:]:
        if _YEAR_RE.search(text) and len(text.split()) <= 8:
            meta = meta or _model(p)
        elif detail is None and meta is None:
            detail = _model(p)
        if meta and detail:
            break
    spacer = None
    seen_content = False
    for p, text in zip(paragraphs, texts, strict=True):
        if text:
            seen_content = True
        elif seen_content and not _has_drawing(p):
            spacer = p
            break
    return _SectionModels(title, detail, meta, _model(bullets[0]) if bullets else None, spacer,
                          _model(filled[0][0]) if filled else fallback)


def _list_bullet_model(document: DocumentObject, base: _Model) -> _Model:
    for style in document.styles.element.iter(qn("w:style")):
        name = style.find(qn("w:name"))
        if name is not None and name.get(qn("w:val")) in {"List Bullet", "Liste à puces"}:
            ppr = OxmlElement("w:pPr")
            style_ref = OxmlElement("w:pStyle")
            style_ref.set(qn("w:val"), style.get(qn("w:styleId")))
            ppr.append(style_ref)
            return _Model(ppr, [_clean_rpr(base.first, bold=False)], None, False)
    return _Model(base.ppr, [_clean_rpr(base.first, bold=False)], None, False, "• ")


def _write_entry(writer: _Writer, models: _SectionModels, entry: Entry, bullet: _Model) -> None:
    title_model = models.title
    segments: list[tuple[str, etree._Element | None]] = [(entry.title, title_model.first)]
    subtitle_inline = bool(entry.subtitle) and models.detail is None
    if subtitle_inline:
        segments.append((f" — {entry.subtitle}" if entry.title else entry.subtitle, title_model.second()))
    period_inline = bool(entry.period) and title_model.has_tab
    if period_inline:
        segments.append((f"\t{entry.period}", title_model.tab_rpr if title_model.tab_rpr is not None
                         else title_model.first))
    writer.paragraph(title_model, segments)
    if entry.subtitle and not subtitle_inline:
        writer.paragraph(models.detail, [(entry.subtitle, models.detail.first)])  # type: ignore[union-attr]
    plain = models.meta or models.detail or _Model(title_model.ppr, [_clean_rpr(title_model.first, bold=False)],
                                                   None, False)
    if entry.period and not period_inline:
        writer.paragraph(plain, [(entry.period, plain.first)])
    for detail in [*entry.details, *entry.tech]:
        model = models.detail or plain
        writer.paragraph(model, [(detail, model.first)])
    for text in entry.bullets:
        writer.paragraph(bullet, [(bullet.text_bullet + text, bullet.first)])


def _write_section(document: DocumentObject, writer: _Writer, models: _SectionModels, view: SectionView,
                   language: str) -> None:
    bullet = models.bullet or _list_bullet_model(document, models.title)
    for position, entry in enumerate(view.entries):
        if position and models.spacer is not None:
            writer.clone(models.spacer)
        _write_entry(writer, models, entry, bullet)
    line_model = models.line
    for label, text in view.lines:
        prefix = line_model.text_bullet
        if label and len(line_model.rprs) > 1:
            writer.paragraph(line_model, [(f"{prefix}{label}{colon(language)}", line_model.first),
                                          (text, line_model.second())])
        else:
            text = f"{prefix}{label}{colon(language)}{text}" if label else prefix + text
            writer.paragraph(line_model, [(text, line_model.first)])


def _set_text(p: etree._Element, parent: object, text: str, *, keep_label: bool = False) -> None:
    model = _model(p)
    label = ""
    if keep_label:
        match = _SUBJECT_RE.match(paragraph_text(p))
        label = match.group(0) if match else ""
    if compact_text(paragraph_text(p)).isupper() and len(text) < 60 and not keep_label:
        text = text.upper()
    _clear_text(p)
    if label:
        _append_run(p, parent, label, model.first)
        _append_run(p, parent, text, model.second() if len(model.rprs) > 1 else model.first)
    else:
        _append_run(p, parent, text, model.first)


def _remove(p: etree._Element) -> None:
    if _has_drawing(p):
        _clear_text(p)
        return
    parent = p.getparent()
    if parent is None:
        return
    # Une cellule de tableau ou une zone de texte doit garder au moins un paragraphe.
    if parent.tag in {qn("w:tc"), _TXBX} and len(parent.findall(qn("w:p"))) <= 1:
        _clear_text(p)
        return
    parent.remove(p)


def _drop_textbox_fallbacks(root: etree._Element) -> None:
    """Supprime les copies VML de secours des zones de texte : elles garderaient l'ancien contenu."""
    for alternate in list(root.iter(_MC_ALTERNATE)):
        if any(True for _ in alternate.iter(_TXBX)):
            for fallback in alternate.findall(_MC_FALLBACK):
                alternate.remove(fallback)


def _scrub(document: DocumentObject, title: str, author: str) -> None:
    properties = document.core_properties
    properties.title = title
    properties.author = author
    properties.last_modified_by = author
    properties.comments = ""
    settings = document.settings.element
    for node in settings.findall(qn("w:attachedTemplate")):
        settings.remove(node)


def _resolve(document: DocumentObject, template_map: TemplateMap) -> dict[str, tuple[list[etree._Element], object]]:
    resolved: dict[str, tuple[list[etree._Element], object]] = {}
    for part in ("body", "header", "first_page_header"):
        root = story_root(document, part)
        if root is not None:
            resolved[part] = (story_paragraphs(root[0]), root[1])
    body = resolved["body"][0]
    for section in template_map.sections:
        start, end = int(section.get("start", -1)), int(section.get("end", -1))
        if section.get("part", "body") != "body" or not (0 <= start < end <= len(body)):
            raise RenderError("La cartographie ne correspond plus au modèle DOCX : relance l'analyse du modèle.")
        if section.get("type") in CV_SECTION_TYPES and identify_section(paragraph_text(body[start])) is None:
            raise RenderError("La cartographie ne correspond plus au modèle DOCX : relance l'analyse du modèle.")
    return resolved


def _ref(resolved: dict, ref: object) -> tuple[etree._Element, object] | None:
    if not isinstance(ref, dict) or ref.get("part") not in resolved:
        return None
    paragraphs, parent = resolved[ref["part"]]
    index = ref.get("index")
    if not isinstance(index, int) or not 0 <= index < len(paragraphs):
        raise RenderError("La cartographie ne correspond plus au modèle DOCX : relance l'analyse du modèle.")
    return paragraphs[index], parent


def _fill_identity(document: DocumentObject, template_map: TemplateMap, resolved: dict, identity: dict,
                   language: str, *, headline: str | None = None) -> list[str]:
    warnings: list[str] = []
    name, contacts, links = identity_lines(identity, language)
    header = template_map.header
    name_ref = _ref(resolved, header.get("name"))
    contact_refs = [ref for ref in (_ref(resolved, item) for item in header.get("contacts", [])) if ref]
    lines = [line for line in (contacts, links) if line]
    if name_ref is not None and name:
        _set_text(name_ref[0], name_ref[1], name)
    elif name:
        lines.insert(0, [name])
    headline_ref = _ref(resolved, header.get("headline"))
    if headline_ref is not None and headline:
        _set_text(headline_ref[0], headline_ref[1], headline)
    if contact_refs:
        if len(contact_refs) == 1 and len(lines) > 1:
            # Une seule ligne de coordonnées dans le modèle : les liens vont sur une copie de cette ligne.
            p, parent = contact_refs[0]
            for _ in lines[1:]:
                copy = deepcopy(p)
                p.addnext(copy)
                contact_refs.append((copy, parent))
                p = copy
            groups = [" · ".join(line) for line in lines]
        else:
            items = [item for line in lines for item in line]
            size = -(-len(items) // len(contact_refs)) if items else 0
            groups = [" · ".join(items[i * size:(i + 1) * size]) for i in range(len(contact_refs))]
        for (p, parent), text in zip(contact_refs, groups, strict=False):
            if text:
                _set_text(p, parent, text)
            else:
                _remove(p)
    elif lines:
        body, parent = resolved["body"]
        anchor = name_ref[0] if name_ref is not None and name_ref[1] is parent else None
        for line in reversed(lines):
            p = OxmlElement("w:p")
            if anchor is not None:
                anchor.addnext(p)
            elif body:
                body[0].addprevious(p)
            else:
                document.element.body.insert(0, p)
            _append_run(p, parent, " · ".join(line), None)
        warnings.append("Le modèle d'origine n'avait pas de ligne de coordonnées repérable : elles ont été ajoutées "
                        "en tête du document.")
    return warnings


def _block(heading: etree._Element, boundaries: set[etree._Element]) -> list[etree._Element]:
    """Titre de section et paragraphes qui le suivent jusqu'au titre suivant (ou à un tableau, une fin de conteneur)."""
    block = [heading]
    for sibling in heading.itersiblings():
        if sibling in boundaries or sibling.tag != qn("w:p"):
            break
        block.append(sibling)
    return block


def fill_cv(document: DocumentObject, template_map: TemplateMap, composition: dict, identity: dict) -> list[str]:
    """Remplit le DOCX d'origine en place ; renvoie les avertissements (sans donnée personnelle).

    Les sections gardent leurs titres, styles et paragraphes d'exemple. Quand toutes sont dans le même
    conteneur (CV à une colonne), elles sont replacées dans l'ordre de la composition, à l'emplacement de la
    première ; sinon (tableau, zones de texte) chacune reste à sa place.
    """
    language = language_of(composition)
    resolved = _resolve(document, template_map)
    body, parent = resolved["body"]
    texts = [compact_text(paragraph_text(p)) for p in body]
    warnings: list[str] = []
    mapped = [section for section in template_map.sections if section.get("type") in (*CV_SECTION_TYPES, "summary")]
    if not any(section["type"] in CV_SECTION_TYPES for section in mapped):
        raise RenderError("Aucune section du modèle DOCX d'origine n'a été reconnue : choisis un modèle JobRadar.")
    boundaries = {body[int(section["start"])] for section in template_map.sections if section.get("part", "body")
                  == "body" and section.get("type") not in LETTER_TYPES}

    captured: dict[str, tuple[etree._Element, list[etree._Element], _SectionModels]] = {}
    duplicates: list[etree._Element] = []
    neutral = _Model(None, [None], None, False)
    for section in mapped:
        start, end = int(section["start"]), int(section["end"])
        if section["type"] in captured:
            # Une composition regroupe tous les éléments d'un type : une seconde rubrique du même type est retirée.
            duplicates.extend(body[start:end])
            continue
        content = body[start + 1:end]
        captured[section["type"]] = (body[start], content,
                                     _section_models(document, content, texts[start + 1:end], neutral))
    cv_kinds = [kind for kind in captured if kind in CV_SECTION_TYPES]
    reference_heading = captured[cv_kinds[0]][0]
    uppercase_headings = all(compact_text(paragraph_text(captured[kind][0])).isupper() for kind in cv_kinds)
    containers = [container_of(captured[kind][0]) for kind in cv_kinds]
    linear = all(container is containers[0] for container in containers)

    def replace_content(content: list[etree._Element], heading: etree._Element) -> _Writer:
        filled = [p for p in content if compact_text(paragraph_text(p))]
        anchor = heading
        if filled:
            first_index = content.index(filled[0])
            last_index = content.index(filled[-1])
            anchor = content[first_index - 1] if first_index else heading
            for p in content[first_index:last_index + 1]:
                _remove(p)
        return _Writer(parent, anchor)

    views = section_views(composition)
    wanted = {view.type for view in views}
    blocks: list[list[etree._Element]] = []
    previous: etree._Element | None = None
    first_heading = min(body.index(captured[kind][0]) for kind in cv_kinds)
    marker = OxmlElement("w:p")  # emplacement de la première section du modèle
    body[first_heading].addprevious(marker)
    for view in views:
        if view.type in captured:
            heading, content, models = captured[view.type]
            writer = replace_content(content, heading)
            _write_section(document, writer, models, view, language)
        else:
            family = {"project": "experience", "experience": "project", "certification": "education"}.get(view.type)
            source = captured.get(family or "") or captured.get("experience") or captured[cv_kinds[0]]
            heading = deepcopy(reference_heading)
            if previous is not None:
                previous.addnext(heading)
            else:
                body[first_heading].addprevious(heading)
            boundaries.add(heading)
            title = section_title(view.type, language)
            _set_text(heading, parent, title.upper() if uppercase_headings else title)
            writer = _Writer(parent, heading)
            _write_section(document, writer, source[2], view, language)
        previous = writer.anchor
        blocks.append([heading])

    for kind in cv_kinds:
        if kind not in wanted:
            heading, content, _ = captured[kind]
            for p in [heading, *content]:
                _remove(p)
    for p in duplicates:
        _remove(p)

    if linear:
        complete = [_block(block[0], boundaries | {marker}) for block in blocks]
        for block in complete:
            for element in block:
                element.getparent().remove(element)
        for block in complete:
            for element in block:
                marker.addprevious(element)
    marker.getparent().remove(marker)

    summary = compact_text(composition.get("summary"))
    if "summary" in captured:
        heading, content, models = captured["summary"]
        if summary:
            writer = replace_content(content, heading)
            writer.paragraph(models.line, [(summary, models.line.first)])
        else:
            for p in [heading, *content]:
                _remove(p)
    else:
        summary_ref = _ref(resolved, template_map.header.get("summary"))
        if summary_ref is not None:
            if summary:
                _set_text(summary_ref[0], summary_ref[1], summary)
            else:
                _remove(summary_ref[0])
        elif summary:
            # Pas d'emplacement prévu : le résumé est placé juste avant la première section, au style du texte courant.
            first = blocks[0][0] if blocks else reference_heading
            p = OxmlElement("w:p")
            first.addprevious(p)
            _append_run(p, parent, summary, None)

    warnings += _fill_identity(document, template_map, resolved, identity, language,
                               headline=compact_text(composition.get("headline")))
    for part_root in (story_root(document, part) for part in ("body", "header", "first_page_header")):
        if part_root is not None:
            _drop_textbox_fallbacks(part_root[0])
    return warnings


def fill_letter(document: DocumentObject, template_map: TemplateMap, letter: dict, identity: dict) -> list[str]:
    language = language_of(letter)
    resolved = _resolve(document, template_map)
    body, parent = resolved["body"]
    parts = {section["type"]: section for section in template_map.sections if section.get("type") in LETTER_TYPES}
    if "letter_greeting" not in parts or "letter_closing" not in parts:
        raise RenderError("La structure de la lettre d'origine (formule d'appel, formule de politesse) n'a pas été "
                          "reconnue : choisis un modèle JobRadar.")
    warnings: list[str] = []
    elements = {kind: body[int(section["start"]):int(section["end"])] for kind, section in parts.items()}
    name, _ = identity_parts(identity, language)

    if "letter_date" in elements:
        _set_text(elements["letter_date"][0], parent, letter_date(language, compact_text(identity.get("city"))))
    if "letter_recipient" in elements:
        lines = letter_recipient(letter)
        existing = [p for p in elements["letter_recipient"] if compact_text(paragraph_text(p))]
        for position, p in enumerate(existing):
            if position < len(lines):
                _set_text(p, parent, lines[position])
            else:
                _remove(p)
        if existing and len(lines) > len(existing):
            # Lignes supplémentaires : clones du dernier paragraphe du bloc, pour garder sa mise en forme.
            writer = _Writer(parent, existing[-1])
            for line in lines[len(existing):]:
                writer.clone(existing[-1])
                _set_text(writer.anchor, parent, line)
        if not lines:
            warnings.append("Le bloc destinataire de la lettre d'origine a été retiré (il visait une autre "
                            "candidature) : ajoute-le dans le DOCX si besoin.")
    subject = compact_text(letter.get("subject"))
    if "letter_subject" in elements:
        if subject:
            _set_text(elements["letter_subject"][0], parent, subject, keep_label=True)
        else:
            _remove(elements["letter_subject"][0])
    elif subject:
        greeting = elements["letter_greeting"][0]
        model = _model(greeting)
        writer = _Writer(parent, greeting.getprevious() if greeting.getprevious() is not None else greeting)
        label = "Subject: " if language == "en" else "Objet : "
        writer.paragraph(model, [(label + subject, _clean_rpr(model.first, bold=True))])
    if compact_text(letter.get("greeting")):
        _set_text(elements["letter_greeting"][0], parent, compact_text(letter.get("greeting")))
    paragraphs = [compact_text(text) for text in letter.get("paragraphs", []) if compact_text(text)] \
        if isinstance(letter.get("paragraphs"), list) else []
    if "letter_body" in elements:
        old = elements["letter_body"]
        filled = [p for p in old if compact_text(paragraph_text(p))]
        model = _model(filled[0])
        spacer = next((p for p in old if not compact_text(paragraph_text(p)) and not _has_drawing(p)), None)
        before = old[0].getprevious()
        writer = _Writer(parent, before if before is not None else elements["letter_greeting"][0])
        for p in old:
            _remove(p)
        for position, text in enumerate(paragraphs):
            if position and spacer is not None:
                writer.clone(spacer)
            writer.paragraph(model, [(text, model.first)])
    else:
        writer = _Writer(parent, elements["letter_greeting"][0])
        model = _Model(_clean_ppr(elements["letter_greeting"][0].find(qn("w:pPr"))),
                       [_model(elements["letter_greeting"][0]).first], None, False)
        for text in paragraphs:
            writer.paragraph(model, [(text, model.first)])
    if compact_text(letter.get("closing")):
        _set_text(elements["letter_closing"][0], parent, compact_text(letter.get("closing")))
    if "letter_signature" in elements and name:
        _set_text(elements["letter_signature"][0], parent, name)
    warnings += _fill_identity(document, template_map, resolved, identity, language)
    for part_root in (story_root(document, part) for part in ("body", "header", "first_page_header")):
        if part_root is not None:
            _drop_textbox_fallbacks(part_root[0])
    return warnings


def open_template(template_map: TemplateMap) -> DocumentObject:
    try:
        return Document(str(template_map.source_path))
    except Exception as exc:
        raise RenderError("Le modèle DOCX d'origine ne peut pas être rouvert.") from exc


def finish(document: DocumentObject, title: str, author: str) -> None:
    _scrub(document, title, author)
