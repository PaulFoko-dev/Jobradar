"""Modèle d'origine DOCX : cartographie, risques ATS, remplissage fidèle et conversion PDF avec ou sans LibreOffice."""

from __future__ import annotations

import importlib.util
import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor
from pypdf import PdfReader

from jobradar.render import RenderError, TemplateMap, analyze_docx_template, ats_check, render_cv, render_letter
from jobradar.render.template import paragraph_text, story_paragraphs

FIXTURES = Path(__file__).parent / "fixtures" / "render"
_SPEC = importlib.util.spec_from_file_location("modeles_docx", FIXTURES / "modeles_docx.py")
assert _SPEC is not None and _SPEC.loader is not None
modeles = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(modeles)

ORIGINAL = {"kind": "original", "id": "src-1"}


def _load(name: str) -> dict:
    path = FIXTURES / "compositions" / f"{name}.json"
    return json.loads((path if path.exists() else FIXTURES / f"{name}.json").read_text(encoding="utf-8"))


@pytest.fixture(autouse=True)
def no_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("jobradar.render.renderer.shutil.which", lambda _name: None)


def _fill(tmp_path: Path, builder, composition_name: str = "stage_cyber"):
    source = builder(tmp_path / "modele.docx")
    mapping = analyze_docx_template(source)
    composition, identity = _load(composition_name), _load("identity")
    render = render_letter if composition["kind"] == "lettre" else render_cv
    result = render(composition, identity, template=ORIGINAL, template_map=mapping, out_dir=tmp_path / "sortie",
                    base_name="Camille_Martin_Original")
    report = ats_check(result.pdf_path, composition=composition, identity=identity, offer_keywords=[],
                       owned_keywords=[], docx_path=result.docx_path)
    return mapping, result, report, Document(result.docx_path)


def _body_texts(document) -> list[str]:
    return [paragraph_text(p) for p in story_paragraphs(document.element.body)]


def _paragraph(document, startswith: str):
    return next(p for p in document.paragraphs if p.text.startswith(startswith))


def test_simple_template_is_mapped_and_compatible(tmp_path: Path) -> None:
    mapping = analyze_docx_template(modeles.simple_ats(tmp_path / "modele.docx"))
    assert [(s["type"], s["title"]) for s in mapping.sections] == [
        ("summary", "Profil"),
        ("experience", "Expériences professionnelles"),
        ("education", "Formation"),
        ("skills", "Compétences"),
        ("language", "Langues"),
        ("other", "Centres d'intérêt"),
    ]
    assert mapping.sections[1]["style"]["heading"]["paragraph_style"] == "Heading 1"
    assert mapping.sections[1]["style"]["content"]["font"]["name"] == "Georgia"
    assert mapping.header == {"location": "body", "name": {"part": "body", "index": 0},
                              "headline": {"part": "body", "index": 1}, "summary": None,
                              "contacts": [{"part": "body", "index": 2}]}
    assert mapping.issues == [] and mapping.ats_compatible is True
    assert TemplateMap.from_json(json.loads(json.dumps(mapping.to_json()))) == mapping


def test_simple_template_is_filled_without_losing_formatting(tmp_path: Path) -> None:
    _, result, report, document = _fill(tmp_path, modeles.simple_ats)
    texts = _body_texts(document)
    assert not any("ncien" in text for text in texts)  # aucun reste du CV de base
    assert "Course à pied, photographie argentique." in texts  # section inconnue conservée telle quelle
    assert document.paragraphs[0].text == "Camille Martin" and document.paragraphs[0].style.name == "Title"
    assert document.paragraphs[1].text == "Stagiaire Analyste SOC"
    assert document.paragraphs[1].runs[0].font.color.rgb == modeles.GREY
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Profil", "Expériences professionnelles", "Projets", "Formation", "Certifications",
                        "Langues", "Compétences", "Centres d'intérêt"]

    title = _paragraph(document, "Analyste SOC (alternance)")
    assert title.text == "Analyste SOC (alternance)\t09/2024 – Aujourd'hui"
    assert title.runs[0].bold and title.runs[0].font.name == modeles.HEADING_FONT
    assert title.runs[1].font.color.rgb == modeles.GREY
    assert title.paragraph_format.tab_stops[0].alignment == 2  # tabulation droite d'origine
    subtitle = _paragraph(document, "Acme Sécurité")
    assert subtitle.runs[0].italic and subtitle.runs[0].font.color.rgb == modeles.GREY
    assert _paragraph(document, "Qualification de 40 alertes").style.name == "List Bullet"
    label = _paragraph(document, "Détection : ")
    assert label.runs[0].bold and label.runs[0].font.color.rgb == modeles.NAVY and not label.runs[1].bold

    assert any("LibreOffice" in warning for warning in result.warnings)
    with zipfile.ZipFile(result.docx_path) as archive:
        xml = b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
    assert str(tmp_path).encode() not in xml and str(tmp_path).encode() not in result.pdf_path.read_bytes()
    assert PdfReader(result.pdf_path).metadata.title == "Camille Martin — Stagiaire Analyste SOC"
    assert report.readability == 100, [issue["message"] for issue in report.issues]


def test_two_column_table_is_flagged_and_filled_in_its_cells(tmp_path: Path) -> None:
    mapping, result, report, document = _fill(tmp_path, modeles.two_columns_table)
    assert any("tableau" in issue for issue in mapping.issues) and mapping.ats_compatible is False
    assert mapping.header["location"] == "table"
    assert {s["container"] for s in mapping.sections} == {"table"}
    left, right = document.tables[0].rows[0].cells
    assert "Détection : Splunk (avancé)" in left.text and "Analyste SOC (alternance)" not in left.text
    assert "Analyste SOC (alternance)" in right.text and "Ancien" not in right.text
    assert any("risques pour les ATS" in warning for warning in result.warnings)
    assert report.readability < 95 and report.ready is False
    assert "layout_columns" in {issue["code"] for issue in report.issues}


def test_text_box_is_flagged_and_its_contacts_are_replaced(tmp_path: Path) -> None:
    mapping, _, report, document = _fill(tmp_path, modeles.text_box)
    assert any("zone(s) de texte" in issue and "coordonnées" in issue for issue in mapping.issues)
    assert mapping.header["location"] == "textbox"
    boxes = list(document.element.body.iter(qn("w:txbxContent")))
    assert len(boxes) == 1  # la copie VML de secours, qui gardait l'ancien contenu, est retirée
    box_text = " ".join(paragraph_text(p) for p in boxes[0].iter(qn("w:p")))
    assert "camille.martin@example.test" in box_text and "ancien@" not in box_text
    assert report.readability < 100  # un ATS qui ignore les zones de texte perd les coordonnées
    assert any(issue["field"] == "identity.email" and "DOCX" in issue["message"] for issue in report.issues)


def test_header_contacts_are_flagged_and_replaced_in_place(tmp_path: Path) -> None:
    mapping, _, report, document = _fill(tmp_path, modeles.header_contacts)
    assert any("en-tête" in issue for issue in mapping.issues) and mapping.ats_compatible is False
    assert mapping.header["location"] == "header"
    header = [p.text for p in document.sections[0].header.paragraphs]
    assert header[0] == "Camille Martin"
    assert header[1].startswith("Paris · camille.martin@example.test")
    assert document.sections[0].header.paragraphs[0].runs[0].font.color.rgb == modeles.NAVY
    assert not any("camille.martin@" in text for text in _body_texts(document))
    assert report.readability < 100
    assert any(issue["field"] == "identity.full_name" for issue in report.issues)


def test_original_letter_is_mapped_and_filled(tmp_path: Path) -> None:
    mapping, result, report, document = _fill(tmp_path, modeles.letter, "lettre_fr")
    assert [s["type"] for s in mapping.sections] == ["letter_recipient", "letter_date", "letter_subject",
                                                     "letter_greeting", "letter_body", "letter_closing",
                                                     "letter_signature"]
    assert mapping.ats_compatible is True
    texts = [text for text in _body_texts(document) if text]
    assert not any("ncien" in text for text in texts)
    subject = _paragraph(document, "Objet : ")
    assert subject.runs[0].bold and subject.text.endswith("(réf. SOC-2027-03)") and not subject.runs[1].bold
    assert texts[1].startswith("Paris · ") and texts[2].startswith("GitHub : ")  # liens sur une copie de la ligne
    assert texts[3].startswith("Paris, le ")
    letter = _load("lettre_fr")
    greeting = texts.index("Madame, Monsieur,")
    assert texts[greeting + 1:greeting + 5] == letter["paragraphs"]
    assert texts[-1] == "Camille Martin"
    assert _paragraph(document, "Actuellement").alignment == 3  # justification d'origine conservée
    assert any("destinataire" in warning for warning in result.warnings)
    assert report.readability == 100, [issue["message"] for issue in report.issues]


def test_letter_with_date_before_recipient_drops_the_old_recipient(tmp_path: Path) -> None:
    mapping, _, report, document = _fill(tmp_path, lambda path: modeles.letter(path, date_first=True), "lettre_fr")
    assert {"letter_date", "letter_recipient"} <= {s["type"] for s in mapping.sections}
    assert not any("ncien" in text or "69000" in text for text in _body_texts(document))
    assert report.readability == 100


def test_letter_recipient_block_is_readdressed_to_the_offer_company(tmp_path: Path) -> None:
    source = modeles.letter(tmp_path / "modele.docx")
    mapping = analyze_docx_template(source)
    letter = _load("lettre_fr") | {"recipient": ["Service recrutement", "Nouvelle Entreprise", "Tour Alpha",
                                                 "92000 Nanterre"]}
    result = render_letter(letter, _load("identity"), template=ORIGINAL, template_map=mapping,
                           out_dir=tmp_path / "sortie", base_name="Camille_Martin_Original")
    texts = [text for text in _body_texts(Document(result.docx_path)) if text]
    start = texts.index("Service recrutement")
    assert texts[start:start + 4] == letter["recipient"]
    assert not any("ncien" in text or "69000" in text for text in texts)
    assert not any("destinataire" in warning for warning in result.warnings)


def test_page_number_footer_is_not_an_ats_issue(tmp_path: Path) -> None:
    mapping = analyze_docx_template(modeles.page_number_footer(tmp_path / "modele.docx"))
    assert mapping.issues == [] and mapping.ats_compatible is True


def test_libreoffice_is_used_when_installed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    if not sys.platform.startswith("linux"):
        pytest.skip("faux exécutable shell")
    reference = tmp_path / "reference.pdf"
    source = modeles.simple_ats(tmp_path / "modele.docx")
    mapping = analyze_docx_template(source)
    composition, identity = _load("stage_cyber"), _load("identity")
    first = render_cv(composition, identity, template=ORIGINAL, template_map=mapping, out_dir=tmp_path / "a",
                      base_name="cv")
    reference.write_bytes(first.pdf_path.read_bytes())
    fake = tmp_path / "soffice"
    log = tmp_path / "appel.txt"
    fake.write_text(
        "#!/bin/sh\n"
        f"echo \"$@\" > '{log}'\n"
        "while [ $# -gt 1 ]; do if [ \"$1\" = --outdir ]; then out=$2; fi; shift; done\n"
        f"cp '{reference}' \"$out/document.pdf\"\n"
    )
    fake.chmod(0o755)
    monkeypatch.setattr("jobradar.render.renderer.shutil.which",
                        lambda name: str(fake) if name == "soffice" else None)
    result = render_cv(composition, identity, template=ORIGINAL, template_map=mapping, out_dir=tmp_path / "b",
                       base_name="cv")
    assert not any("LibreOffice" in warning for warning in result.warnings)
    arguments = log.read_text()
    assert "--headless" in arguments and "--convert-to pdf" in arguments and "UserInstallation=file://" in arguments
    assert PdfReader(result.pdf_path).metadata.author == "Camille Martin"


def test_stale_or_unreadable_templates_are_controlled_errors(tmp_path: Path) -> None:
    source = modeles.simple_ats(tmp_path / "modele.docx")
    mapping = analyze_docx_template(source)
    document = Document(source)
    for paragraph in document.paragraphs[:6]:
        paragraph._p.getparent().remove(paragraph._p)
    document.save(source)
    with pytest.raises(RenderError, match="cartographie"):
        render_cv(_load("stage_cyber"), _load("identity"), template=ORIGINAL, template_map=mapping,
                  out_dir=tmp_path / "sortie", base_name="cv")
    broken = tmp_path / "faux.docx"
    broken.write_bytes(b"pas un docx")
    with pytest.raises(RenderError, match="ne peut pas être lu"):
        analyze_docx_template(broken)
    with pytest.raises(RenderError):
        analyze_docx_template(tmp_path / "absent.docx")


def test_document_without_sections_is_not_ats_compatible(tmp_path: Path) -> None:
    path = tmp_path / "libre.docx"
    document = Document()
    document.add_paragraph("Prénom Nom")
    document.add_paragraph("Un CV rédigé comme une lettre, sans aucun titre de section.")
    document.save(path)
    mapping = analyze_docx_template(path)
    assert mapping.ats_compatible is False
    assert any("titre de section" in issue for issue in mapping.issues)
    with pytest.raises(RenderError, match="modèle JobRadar"):
        render_cv(_load("stage_cyber"), _load("identity"), template=ORIGINAL, template_map=mapping,
                  out_dir=tmp_path / "sortie", base_name="cv")


def test_example_color_and_font_survive_for_every_cloned_line(tmp_path: Path) -> None:
    _, _, _, document = _fill(tmp_path, modeles.simple_ats, "profil_fourni")
    titles = [p for p in document.paragraphs if "\t" in p.text]
    assert len(titles) >= 9  # 4 expériences, 3 projets, 2 formations (certifications : date en tabulation)
    for paragraph in titles:
        assert paragraph.runs[0].font.name == modeles.HEADING_FONT
        assert paragraph.runs[-1].font.color.rgb == RGBColor(0x59, 0x59, 0x59)
