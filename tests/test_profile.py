"""Import et cycle de vie du profil candidat, sans vraie IA."""
import io
import json
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from jobradar import ai, profile

IDENTITY = ("Alice Martin", "alice.martin@example.test", "06 12 34 56 78", "Paris", "github.com/alice")


def extraction(**overrides):
    return {
        "headline": "Analyste SOC", "summary": "Profil cyber",
        "experiences": [], "projects": [], "education": [], "certifications": [],
        "languages": [], "skills": [],
    } | overrides


def experience(**overrides):
    return {"title": "Analyste SOC", "organization": "Acme", "location": "Lyon", "start": "2024-09",
            "end": "2025-03", "context": "Stage", "technologies": ["Splunk"],
            "bullets": [{"text": "Analyse de 20 alertes par jour"}]} | overrides


def text_pdf(text: str) -> bytes:
    """PDF textuel minimal (une page, Helvetica) avec une table de références exacte."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    output += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(output)


@pytest.fixture
def prompts(data_dir, monkeypatch):
    """Faux assistant : enregistre chaque prompt et renvoie les structurations prévues, dans l'ordre."""
    class Sent(list):
        answers: list

    sent, answers = Sent(), []

    def fake_ask(prompt, schema, task):
        sent.append({"prompt": prompt, "schema": schema, "task": task})
        return answers.pop(0) if answers else extraction()

    monkeypatch.setattr(ai, "ask_ai", fake_ask)
    sent.answers = answers
    return sent


def test_import_text_extracts_identity_locally_and_never_sends_it_to_ai(prompts):
    prompts.answers.append(extraction(skills=[{"name": "Python", "category": "Langages", "level": "avance"}]))
    content = ("Alice Martin\nalice.martin@example.test\n06 12 34 56 78\n75001 Paris\n"
               "https://github.com/alice\nCompétences\nPython avancé").encode()

    result = profile.import_source(kind="cv", name="CV SOC", tags=["soc"], filename="cv.txt", content=content)

    assert result["profile"]["identity"] == {
        "full_name": "Alice Martin", "email": "alice.martin@example.test", "phone": "06 12 34 56 78",
        "address_line1": "", "postal_code": "", "city": "Paris", "region": "", "country": "",
        "links": [{"label": "GitHub", "url": "https://github.com/alice"}],
    }
    prompt = prompts[0]["prompt"]
    assert prompts[0]["task"] == "profile"
    for secret in (*IDENTITY, "alice"):
        assert secret.casefold() not in prompt.casefold()
    for marker in ("{{NOM}}", "{{EMAIL}}", "{{TELEPHONE}}", "{{VILLE}}", "{{LIENS}}"):
        assert marker in prompt
    assert (profile.sources_dir() / "src-1.txt").read_bytes() == content
    assert result["source"]["id"] == "src-1" and result["source"]["language"] == "fr"


def test_identity_known_from_a_first_document_is_masked_everywhere_in_the_next(prompts):
    profile.import_source(kind="cv", name="CV", tags=[], filename="cv.txt",
                          content=b"Alice Martin\nalice.martin@example.test\n06 12 34 56 78\n75001 Paris")
    letter = ("Madame, Monsieur,\nJe vis à Paris et suis joignable au 06.12.34.56.78 ou +33 6 12 34 56 78.\n"
              "Profil : linkedin.com/in/alice-martin — autre adresse : a.martin@perso.test\n"
              "Cordialement,\nALICE MARTIN\nMartin Alice").encode()

    profile.import_source(kind="lettre", name="Lettre", tags=[], filename="lettre.md", content=letter)

    prompt = prompts[1]["prompt"]
    for secret in ("Alice", "Martin", "Paris", "06.12.34.56.78", "6 12 34 56 78", "linkedin.com/in",
                   "a.martin@perso.test"):
        assert secret.casefold() not in prompt.casefold(), secret
    assert "Cordialement" in prompt and "Madame, Monsieur" in prompt


def test_docx_header_table_identity_is_masked(prompts):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Alice Martin"
    table.cell(0, 1).text = "alice.martin@example.test | 06 12 34 56 78 | Paris"
    document.add_paragraph("Expérience : Analyste SOC chez Acme")
    stream = io.BytesIO()
    document.save(stream)

    result = profile.import_source(kind="cv", name="CV", tags=[], filename="cv.docx", content=stream.getvalue())

    assert result["profile"]["identity"]["city"] == "Paris"
    for secret in IDENTITY[:4]:
        assert secret not in prompts[0]["prompt"]
    assert "Analyste SOC chez Acme" in prompts[0]["prompt"]


def test_markers_returned_by_ai_are_restored_locally(prompts):
    prompts.answers.append(extraction(experiences=[experience(location="{{VILLE}}")]))

    result = profile.import_source(kind="cv", name="CV", tags=[], filename="cv.txt",
                                   content="Alice Martin\n75001 Paris\nAnalyste SOC à Paris".encode())

    assert result["profile"]["experiences"][0]["location"] == "Paris"


def test_docx_extraction_keeps_paragraphs_and_tables_in_document_order():
    document = Document()
    document.add_paragraph("Avant")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Poste"
    table.cell(0, 1).text = "Acme"
    document.add_paragraph("Après")
    stream = io.BytesIO()
    document.save(stream)

    assert profile.extract_text("cv.docx", stream.getvalue()).splitlines() == ["Avant", "Poste | Acme", "Après"]


def test_text_pdf_is_extracted():
    assert "Analyste SOC chez Acme" in profile.extract_text("cv.pdf", text_pdf("Analyste SOC chez Acme"))


def test_scanned_pdf_is_rejected_with_actionable_message():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = io.BytesIO()
    writer.write(stream)

    with pytest.raises(profile.ProfileError, match="DOCX ou un fichier texte"):
        profile.extract_text("scan.pdf", stream.getvalue())


@pytest.mark.parametrize(
    ("filename", "content", "message"),
    [("faux.pdf", b"pas un pdf", "ne correspond pas"),
     ("faux.docx", b"%PDF-1.4", "ne correspond pas"),
     ("faux.txt", b"PK\x03\x04archive", "ne correspond pas"),
     ("binaire.md", b"texte\x00binaire", "ne correspond pas"),
     ("faux.exe", b"MZ", "Formats acceptés"),
     ("../cv.txt", b"texte", "nom du fichier"),
     ("vide.txt", b"", "vide")],
)
def test_source_signature_extension_and_name_are_checked(filename, content, message):
    with pytest.raises(profile.ProfileError, match=message):
        profile.extract_text(filename, content)


def test_source_size_and_docx_expansion_are_bounded(monkeypatch):
    with pytest.raises(profile.ProfileError, match="10 Mo"):
        profile.extract_text("cv.txt", b"a" * (profile.MAX_SOURCE_BYTES + 1))
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("word/document.xml", "0" * 5000)
    monkeypatch.setattr(profile, "MAX_DOCX_UNCOMPRESSED", 1000)
    with pytest.raises(profile.ProfileError, match="volumineux"):
        profile.extract_text("cv.docx", stream.getvalue())


def test_mask_identity_covers_phone_variants_emails_and_bare_links():
    identity = {"full_name": "Alice Martin", "email": "alice@example.test", "phone": "06 12 34 56 78",
                "city": "Lyon", "links": [{"label": "GitHub", "url": "https://github.com/alice"}]}
    text = ("MARTIN Alice, 0612345678, +33 6 12 34 56 78, alice@example.test, autre@x.fr, "
            "github.com/alice, https://exemple.test/portfolio, Lyon")

    masked = profile.mask_identity(text, identity)

    for secret in ("MARTIN", "Alice", "0612345678", "6 12 34 56 78", "@", "github", "exemple.test", "Lyon"):
        assert secret not in masked, secret


def test_merge_deduplicates_experience_preserves_variant_and_lists_conflict(data_dir):
    current = profile.merge_extraction(profile.empty_profile(), extraction(experiences=[experience()]),
                                       "src-1", "2026-01-01T00:00:00Z")
    second = extraction(experiences=[experience(
        title="Analyste cybersécurité", technologies=["Sentinel"],
        bullets=[{"text": "Analyse de 20 alertes chaque jour"}, {"text": "Rédaction de procédures"}])])

    merged = profile.merge_extraction(current, second, "src-2", "2026-02-01T00:00:00Z")

    assert len(merged["experiences"]) == 1
    item = merged["experiences"][0]
    assert item["bullets"][0]["variants"] == [{"text": "Analyse de 20 alertes chaque jour", "theme": "",
                                               "validated_at": None,
                                               "provenance": {"source": "src-2", "at": "2026-02-01T00:00:00Z"}}]
    assert [bullet["text"] for bullet in item["bullets"]] == ["Analyse de 20 alertes par jour",
                                                              "Rédaction de procédures"]
    assert item["technologies"] == ["Splunk", "Sentinel"]
    assert merged["conflicts"] == [{"id": "conflict-1", "item_id": item["id"], "field": "title",
                                    "values": ["Analyste SOC", "Analyste cybersécurité"],
                                    "sources": ["src-1", "src-2"], "resolved": False}]


def test_resolved_conflict_is_not_raised_again_and_choice_prevails(prompts):
    prompts.answers += [extraction(experiences=[experience()]),
                        extraction(experiences=[experience(end="2025-06")]),
                        extraction(experiences=[experience(end="2025-06")])]
    for number in (1, 2):
        profile.import_source(kind="cv", name=f"CV {number}", tags=[], filename="cv.txt", content=b"Texte")
    conflict = profile.list_conflicts()[0]
    assert conflict["field"] == "end" and conflict["values"] == ["2025-03", "2025-06"]

    with pytest.raises(profile.ProfileError, match="propositions"):
        profile.resolve_conflict(conflict["id"], "2030-01")
    resolved = profile.resolve_conflict(conflict["id"], "2025-06")
    profile.import_source(kind="cv", name="CV 3", tags=[], filename="cv.txt", content=b"Texte")

    assert resolved["experiences"][0]["end"] == "2025-06"
    assert profile.list_conflicts() == []
    assert len(profile.load_profile()["conflicts"]) == 1


def test_ai_output_is_normalised_dates_levels_and_settings(prompts):
    prompts.answers.append(extraction(
        experiences=[experience(start="09/2024", end="n'importe quoi")],
        projects=[{"title": "CTF", "setting": "inconnu", "start": None, "end": None, "technologies": [],
                   "bullets": [], "link": None}],
        skills=[{"name": "Python", "category": "Langages", "level": "expert"}]))

    saved = profile.import_source(kind="cv", name="CV", tags=[], filename="cv.txt", content=b"Texte")["profile"]

    assert (saved["experiences"][0]["start"], saved["experiences"][0]["end"]) == ("2024-09", None)
    assert saved["projects"][0]["setting"] == "autre"
    assert saved["skills"][0]["level"] == "" and saved["skills"][0]["evidence"] == []


def test_identifiers_are_never_reused_after_deletion(prompts):
    prompts.answers += [extraction(experiences=[experience()]),
                        extraction(experiences=[experience(organization="Autre")])]
    first = profile.import_source(kind="cv", name="CV", tags=[], filename="cv.txt", content=b"A")["profile"]
    old_id = first["experiences"][0]["id"]
    profile.update_profile({"experiences": []})

    second = profile.import_source(kind="cv", name="CV 2", tags=[], filename="cv.txt", content=b"B")["profile"]

    assert old_id == "exp-1"
    assert second["experiences"][0]["id"] == "exp-2"
    assert second["experiences"][0]["bullets"][0]["id"] == "exp-2-b1"
    assert [source["id"] for source in second["sources"]] == ["src-1", "src-2"]


def test_profile_versions_and_restore_are_monotonic(data_dir):
    first = profile.save_profile(profile.empty_profile() | {"headline": "Version une"})
    second = profile.update_profile({"headline": "Version deux"})

    restored = profile.restore_version(first["version"])

    assert second["version"] == 2
    assert restored["version"] == 3
    assert restored["headline"] == "Version une"
    assert {item["version"] for item in profile.list_versions()} == {1, 2, 3}
    stored = json.loads((profile.versions_dir() / "profil-v2.json").read_text(encoding="utf-8"))
    assert stored["headline"] == "Version deux"
    assert not list(profile.profile_dir().rglob("*.tmp"))


def test_update_profile_marks_user_items_and_protects_sources(data_dir):
    profile.save_profile(profile.empty_profile())

    saved = profile.update_profile({"projects": [{"id": "proj-9", "title": "Lab", "bullets": [
        {"id": "proj-9-b1", "text": "Mise en place d'un SIEM"}]}]})

    assert saved["projects"][0]["provenance"]["source"] == "ajout_utilisateur"
    assert saved["projects"][0]["bullets"][0]["provenance"]["source"] == "ajout_utilisateur"
    with pytest.raises(profile.ProfileError, match="sources"):
        profile.update_profile({"sources": [{"id": "src-9"}]})
    with pytest.raises(profile.ProfileError, match="inconnus"):
        profile.update_profile({"inconnu": 1})


def test_last_source_cannot_be_deleted(prompts):
    profile.import_source(kind="cv", name="CV", tags=[], filename="cv.txt", content=b"A")
    profile.import_source(kind="lettre", name="Lettre", tags=[], filename="lettre.txt", content=b"B")

    profile.delete_source("src-1")

    assert not (profile.sources_dir() / "src-1.txt").exists()
    with pytest.raises(profile.ProfileError, match="au moins un"):
        profile.delete_source("src-2")


def test_suggest_base_uses_tags_language_and_keywords(data_dir):
    data = profile.empty_profile()
    data["sources"] = [
        {"id": "src-1", "kind": "cv", "name": "SOC FR", "tags": ["soc"], "language": "fr",
         "keywords": ["Splunk"]},
        {"id": "src-2", "kind": "cv", "name": "Dev EN", "tags": ["dev"], "language": "en",
         "keywords": ["Java"]},
        {"id": "src-3", "kind": "lettre", "name": "Lettre", "tags": [], "language": "fr", "keywords": []},
    ]
    profile.save_profile(data)

    french = profile.suggest_bases({"title": "Stage SOC Splunk", "description": "Analyse des alertes pour le SOC"},
                                   kind="cv")
    english = profile.suggest_bases({"title": "Java developer", "description": "Work with the team and the code"},
                                    kind="cv")

    assert [item["id"] for item in french] == ["src-1", "src-2"]
    assert english[0]["id"] == "src-2"
