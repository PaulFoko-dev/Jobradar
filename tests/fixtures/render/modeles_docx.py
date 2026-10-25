"""Modèles DOCX d'origine construits par python-docx pour les tests du rendu (aucun fichier binaire versionné).

Chaque fonction écrit un CV ou une lettre « d'un utilisateur » avec ses propres polices, couleurs et styles.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADING_FONT = "Georgia"
BODY_FONT = "Verdana"


def _fonts(run_or_style: object, name: str) -> None:
    element = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._r  # type: ignore[attr-defined]
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = parse_xml(f'<w:rFonts xmlns:w="{rpr.nsmap["w"]}"/>')
        rpr.insert(0, fonts)
    for attribute in list(fonts.attrib):
        del fonts.attrib[attribute]
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def _base_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    normal = document.styles["Normal"]
    _fonts(normal, BODY_FONT)
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.1
    heading = document.styles["Heading 1"]
    _fonts(heading, HEADING_FONT)
    heading.font.size = Pt(12)
    heading.font.color.rgb = NAVY
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)
    title = document.styles["Title"]
    _fonts(title, HEADING_FONT)
    title.font.size = Pt(24)
    title.font.color.rgb = NAVY
    title.paragraph_format.space_after = Pt(2)
    return document


def _entry(container: object, title: str, dates: str, subtitle: str, bullets: list[str], width_cm: float) -> None:
    head = container.add_paragraph()  # type: ignore[attr-defined]
    head.paragraph_format.tab_stops.add_tab_stop(Cm(width_cm), WD_TAB_ALIGNMENT.RIGHT)
    head.paragraph_format.space_before = Pt(4)
    run = head.add_run(title)
    run.bold = True
    _fonts(run, HEADING_FONT)
    run.font.size = Pt(10.5)
    dates_run = head.add_run(f"\t{dates}")
    dates_run.font.color.rgb = GREY
    dates_run.font.size = Pt(9)
    sub = container.add_paragraph()  # type: ignore[attr-defined]
    sub_run = sub.add_run(subtitle)
    sub_run.italic = True
    sub_run.font.color.rgb = GREY
    for text in bullets:
        container.add_paragraph(text, style="List Bullet")  # type: ignore[attr-defined]


def _labelled(container: object, label: str, text: str) -> None:
    paragraph = container.add_paragraph()  # type: ignore[attr-defined]
    label_run = paragraph.add_run(f"{label} : ")
    label_run.bold = True
    label_run.font.color.rgb = NAVY
    paragraph.add_run(text)


def _cv_body(container: object, width_cm: float, *, with_profile: bool = True, contacts: bool = True) -> None:
    add = container.add_paragraph  # type: ignore[attr-defined]
    if with_profile:
        add("Profil", style="Heading 1")
        add("Ancienne accroche du CV de base, écrite pour une autre candidature dans un autre domaine.")
    add("Expériences professionnelles", style="Heading 1")
    _entry(container, "Ancien poste exemple", "01/2020 – 06/2021", "Ancienne entreprise · Lyon",
           ["Ancienne réalisation numéro un.", "Ancienne réalisation numéro deux."], width_cm)
    _entry(container, "Autre ancien poste", "09/2018 – 12/2019", "Autre entreprise · Lyon",
           ["Autre ancienne réalisation."], width_cm)
    add("Formation", style="Heading 1")
    _entry(container, "Ancien diplôme", "2016 – 2018", "Ancienne école", [], width_cm)
    add("Compétences", style="Heading 1")
    _labelled(container, "Outils", "ancienne compétence A, ancienne compétence B")
    add("Langues", style="Heading 1")
    add("Anglais : intermédiaire")
    add("Centres d'intérêt", style="Heading 1")
    add("Course à pied, photographie argentique.")


def _identity(container: object, *, headline: bool = True) -> None:
    add = container.add_paragraph  # type: ignore[attr-defined]
    add("Prénom Ancien", style="Title")
    if headline:
        line = add()
        run = line.add_run("Ancien titre de CV")
        run.bold = True
        run.font.color.rgb = GREY
        run.font.size = Pt(11)


def _contact(container: object) -> None:
    paragraph = container.add_paragraph()  # type: ignore[attr-defined]
    run = paragraph.add_run("Lyon · ancien@example.test · 06 00 00 00 00")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY


def simple_ats(path: Path) -> Path:
    """CV classique compatible ATS : une colonne, titres en style « Titre 1 », dates en tabulation droite."""
    document = _base_document()
    _identity(document)
    _contact(document)
    _cv_body(document, 17)
    document.save(path)
    return path


def two_columns_table(path: Path) -> Path:
    """CV « deux colonnes » réalisé par un tableau sans bordure : coordonnées et compétences à gauche."""
    document = _base_document()
    _identity(document)
    table = document.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(5.5), Cm(11.5)
    left.paragraphs[0].add_run("Lyon · ancien@example.test · 06 00 00 00 00").font.size = Pt(8.5)
    left.add_paragraph("Compétences", style="Heading 1")
    _labelled(left, "Outils", "ancienne compétence A")
    left.add_paragraph("Langues", style="Heading 1")
    left.add_paragraph("Anglais : intermédiaire")
    right.paragraphs[0].text = "Expériences professionnelles"
    right.paragraphs[0].style = document.styles["Heading 1"]
    _entry(right, "Ancien poste exemple", "01/2020 – 06/2021", "Ancienne entreprise · Lyon",
           ["Ancienne réalisation numéro un."], 11)
    right.add_paragraph("Formation", style="Heading 1")
    _entry(right, "Ancien diplôme", "2016 – 2018", "Ancienne école", [], 11)
    document.save(path)
    return path


_TEXTBOX = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing>
        <wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="1"
                   behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH>
          <wp:positionV relativeFrom="margin"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="2160000" cy="540000"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapSquare wrapText="bothSides"/>
          <wp:docPr id="10" name="Zone de texte 1"/>
          <wp:cNvGraphicFramePr/>
          <a:graphic>
            <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
              <wps:wsp>
                <wps:cNvSpPr txBox="1"/>
                <wps:spPr>
                  <a:xfrm><a:off x="0" y="0"/><a:ext cx="2160000" cy="540000"/></a:xfrm>
                  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                </wps:spPr>
                <wps:txbx><w:txbxContent>
                  <w:p><w:r><w:rPr><w:color w:val="595959"/><w:sz w:val="17"/></w:rPr>
                    <w:t>ancien@example.test</w:t></w:r></w:p>
                  <w:p><w:r><w:rPr><w:color w:val="595959"/><w:sz w:val="17"/></w:rPr>
                    <w:t>06 00 00 00 00</w:t></w:r></w:p>
                </w:txbxContent></wps:txbx>
                <wps:bodyPr/>
              </wps:wsp>
            </a:graphicData>
          </a:graphic>
        </wp:anchor>
      </w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict><v:shape style="width:170pt;height:42pt"><v:textbox><w:txbxContent>
        <w:p><w:r><w:t>ancien@example.test</w:t></w:r></w:p>
        <w:p><w:r><w:t>06 00 00 00 00</w:t></w:r></w:p>
      </w:txbxContent></v:textbox></v:shape></w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


def text_box(path: Path) -> Path:
    """CV dont les coordonnées sont dans une zone de texte flottante, ancrée dans le paragraphe du nom."""
    document = _base_document()
    _identity(document, headline=False)
    document.paragraphs[0]._p.append(parse_xml(_TEXTBOX))
    _cv_body(document, 17, with_profile=False)
    document.save(path)
    return path


def page_number_footer(path: Path) -> Path:
    """CV simple dont le pied de page ne contient qu'un numéro de page."""
    simple_ats(path)
    document = Document(path)
    footer = document.sections[0].footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = "Page 1 sur 1"
    document.save(path)
    return path


def header_contacts(path: Path) -> Path:
    """CV dont le nom et les coordonnées sont dans l'en-tête de page."""
    document = _base_document()
    header = document.sections[0].header
    header.is_linked_to_previous = False
    name = header.paragraphs[0]
    name_run = name.add_run("Prénom Ancien")
    name_run.bold = True
    name_run.font.size = Pt(16)
    name_run.font.color.rgb = NAVY
    contact = header.add_paragraph()
    contact.add_run("Lyon · ancien@example.test · 06 00 00 00 00").font.size = Pt(8.5)
    _cv_body(document, 17, with_profile=False)
    document.save(path)
    return path


def letter(path: Path, *, date_first: bool = False) -> Path:
    """Lettre de motivation d'origine : coordonnées, destinataire, date, objet, corps et signature."""
    document = _base_document()
    name = document.add_paragraph()
    name_run = name.add_run("Prénom Ancien")
    name_run.bold = True
    name_run.font.size = Pt(14)
    name_run.font.color.rgb = NAVY
    _contact(document)
    document.add_paragraph()
    if date_first:
        document.add_paragraph("Lyon, le 3 mars 2025").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for line in ("Ancienne Société", "Service des ressources humaines", "69000 Lyon"):
        recipient = document.add_paragraph(line)
        recipient.paragraph_format.left_indent = Cm(9)
    if not date_first:
        document.add_paragraph("Lyon, le 3 mars 2025").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()
    subject = document.add_paragraph()
    subject.add_run("Objet : ").bold = True
    subject.add_run("ancienne candidature au poste de technicien")
    document.add_paragraph()
    document.add_paragraph("Madame, Monsieur,")
    document.add_paragraph()
    for text in ("Ancien premier paragraphe de la lettre de base, écrit pour une autre offre.",
                 "Ancien deuxième paragraphe décrivant une autre expérience.",
                 "Ancien troisième paragraphe de conclusion."):
        body = document.add_paragraph(text)
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        document.add_paragraph()
    document.add_paragraph("Je vous prie d'agréer, Madame, Monsieur, mes salutations distinguées.")
    document.add_paragraph()
    document.add_paragraph("Prénom Ancien")
    document.save(path)
    return path
