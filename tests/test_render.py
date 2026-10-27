from __future__ import annotations

import http.server
import json
import os
import sys
import threading
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from jobradar.render import RenderError, ats_check, list_templates, render_cv, render_letter
from jobradar.render.chromium import print_variants, remove_trailing_blank_page
from jobradar.render.jobradar_html import cv_html, letter_html

FIXTURES = Path(__file__).parent / "fixtures" / "render"
COMPOSITIONS = FIXTURES / "compositions"
TEMPLATES = ["sobre-ats", "moderne-ats", "technique-projets"]
# Compositions de référence : (fichier, identité). Stage court, profil fourni (4 expériences, 3 projets), CV anglais.
CV_CASES = [("stage_cyber", "identity"), ("profil_fourni", "identity"), ("cv_anglais", "identity_en")]
LETTER_CASES = [("lettre_fr", "identity"), ("lettre_en", "identity_en")]


def _load(name: str) -> dict:
    path = COMPOSITIONS / f"{name}.json"
    if not path.exists():
        path = FIXTURES / f"{name}.json"
    return json.loads(path.read_text(encoding="utf-8"))


def _render_cv(tmp_path: Path, composition: dict, identity: dict, template_id: str, name: str = "cv"):
    return render_cv(composition, identity, template={"kind": "jobradar", "id": template_id}, template_map=None,
                     out_dir=tmp_path, base_name=name)


def test_list_templates_returns_three_independent_ats_templates() -> None:
    templates = list_templates()
    assert [item["id"] for item in templates] == TEMPLATES
    assert [item["name"] for item in templates] == ["Sobre ATS", "Moderne ATS", "Technique/Projets"]
    assert all(item["ats"] is True and item["description"] for item in templates)
    templates[0]["name"] = "modifié"
    assert list_templates()[0]["name"] == "Sobre ATS"


@pytest.mark.parametrize("template_id", TEMPLATES)
@pytest.mark.parametrize(("case", "identity_name"), CV_CASES)
def test_every_jobradar_cv_is_fully_readable(tmp_path: Path, template_id: str, case: str, identity_name: str) -> None:
    """Test bloquant (étude §3.1) : chaque modèle JobRadar obtient 100 % sur chaque CV de référence."""
    composition, identity = _load(case), _load(identity_name)
    result = _render_cv(tmp_path, composition, identity, template_id, f"Camille_Martin_CV_{case}")

    assert result.pdf_path.name == f"Camille_Martin_CV_{case}.pdf"
    assert result.docx_path is not None and result.docx_path.name == f"Camille_Martin_CV_{case}.docx"
    report = ats_check(result.pdf_path, composition=composition, identity=identity, offer_keywords=[],
                       owned_keywords=[], docx_path=result.docx_path)
    assert report.readability == 100, [issue["message"] for issue in report.issues]
    assert report.ready is True
    assert report.issues == []


@pytest.mark.parametrize("template_id", TEMPLATES)
@pytest.mark.parametrize(("case", "identity_name"), LETTER_CASES)
def test_every_jobradar_letter_is_fully_readable(tmp_path: Path, template_id: str, case: str,
                                                 identity_name: str) -> None:
    letter, identity = _load(case), _load(identity_name)
    result = render_letter(letter, identity, template={"kind": "jobradar", "id": template_id}, template_map=None,
                           out_dir=tmp_path, base_name="lettre")
    assert len(PdfReader(result.pdf_path).pages) == 1
    report = ats_check(result.pdf_path, composition=letter, identity=identity, offer_keywords=[], owned_keywords=[],
                       docx_path=result.docx_path)
    assert report.readability == 100, [issue["message"] for issue in report.issues]
    text = "\n".join(p.text for p in Document(result.docx_path).paragraphs)
    assert letter["paragraphs"][1] in text and letter["closing"] in text


@pytest.mark.parametrize("template_id", TEMPLATES)
def test_internship_cv_fits_one_page_with_metadata_and_embedded_fonts(tmp_path: Path, template_id: str) -> None:
    composition, identity = _load("stage_cyber"), _load("identity")
    result = _render_cv(tmp_path, composition, identity, template_id)
    reader = PdfReader(result.pdf_path)
    assert len(reader.pages) == 1
    assert reader.metadata.title == "Camille Martin — Stagiaire Analyste SOC"
    assert reader.metadata.author == "Camille Martin"
    for page in reader.pages:
        for font in page["/Resources"]["/Font"].values():
            font = font.get_object()
            descendant = font["/DescendantFonts"][0].get_object() if "/DescendantFonts" in font else font
            assert "/FontFile2" in descendant["/FontDescriptor"].get_object()
    assert result.warnings == []


@pytest.mark.parametrize("template_id", TEMPLATES)
def test_long_title_never_has_the_date_inserted_in_it(tmp_path: Path, template_id: str) -> None:
    composition, identity = _load("stage_cyber"), _load("identity")
    composition["sections"][0]["items"][0]["fields"]["title"] = (
        "Analyste en cybersécurité opérationnelle chargé de la détection des menaces et de la réponse aux "
        "incidents (alternance)")
    result = _render_cv(tmp_path, composition, identity, template_id)
    report = ats_check(result.pdf_path, composition=composition, identity=identity, offer_keywords=[],
                       owned_keywords=[], docx_path=result.docx_path)
    assert report.readability == 100, [issue["message"] for issue in report.issues]


def test_rich_profile_stays_within_two_pages(tmp_path: Path) -> None:
    result = _render_cv(tmp_path, _load("profil_fourni"), _load("identity"), "moderne-ats")
    assert len(PdfReader(result.pdf_path).pages) <= 2


def test_cv_follows_ats_writing_rules(tmp_path: Path) -> None:
    composition, identity = _load("stage_cyber"), _load("identity")
    result = _render_cv(tmp_path, composition, identity, "sobre-ats")
    text = PdfReader(result.pdf_path).pages[0].extract_text()
    for title in ("Expérience professionnelle", "Projets", "Formation", "Certifications", "Langues", "Compétences"):
        assert title.upper() in text.upper()
    assert "09/2024 – Aujourd'hui" in text and "06/2023 – 08/2023" in text
    assert "Tél. +33 6 12 34 56 78" in text  # le téléphone est annoncé par un mot, pas par un pictogramme
    assert "KQL (notion)" in text
    english = PdfReader(_render_cv(tmp_path, _load("cv_anglais"), _load("identity_en"), "sobre-ats",
                                   "en").pdf_path).pages[0].extract_text()
    assert "PROFESSIONAL EXPERIENCE" in english.upper() and "09/2024 – Present" in english
    assert "Detection: Splunk" in english  # typographie anglaise : pas d'espace avant le deux-points


def test_docx_is_clean_and_structured(tmp_path: Path) -> None:
    result = _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "technique-projets")
    document = Document(result.docx_path)
    properties = document.core_properties
    assert (properties.title, properties.author, properties.last_modified_by) == (
        "Camille Martin — Stagiaire Analyste SOC", "Camille Martin", "Camille Martin")
    assert properties.comments == ""
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Expérience professionnelle", "Projets", "Formation", "Certifications", "Langues",
                        "Compétences"]
    assert not document.tables and not document.inline_shapes
    assert document.sections[0].header.is_linked_to_previous  # rien dans l'en-tête de page
    with zipfile.ZipFile(result.docx_path) as archive:
        content = b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
    assert b"python-docx" not in content and str(tmp_path).encode() not in content


def test_untrusted_content_is_escaped_everywhere(tmp_path: Path) -> None:
    trap = '<img src="http://127.0.0.1:9/x" onerror="alert(1)"><script>document.title="JS"</script>'
    composition, identity = _load("stage_cyber"), _load("identity")
    composition["summary"] = trap
    composition["headline"] = f"Titre {trap}"
    composition["sections"][0]["items"][0]["bullets"][0]["text"] = trap
    composition["sections"][-1]["groups"][0]["label"] = trap
    identity = dict(identity, full_name=f"Camille {trap}", links=[{"label": trap, "url": "javascript:alert(1)"}])
    html = cv_html(composition, identity, "moderne-ats", title=trap, author=trap)
    assert "<script" not in html and "<img" not in html
    assert "&lt;script&gt;" in html
    letter = dict(_load("lettre_fr"), paragraphs=[trap], subject=trap)
    assert "<script" not in letter_html(letter, identity, "sobre-ats", title=trap, author=trap)

    result = _render_cv(tmp_path, composition, identity, "moderne-ats")
    text = PdfReader(result.pdf_path).pages[0].extract_text()
    assert "<script>" in text  # affiché tel quel, jamais interprété
    docx_text = "\n".join(p.text for p in Document(result.docx_path).paragraphs)
    assert trap in docx_text


class _CountingHandler(http.server.BaseHTTPRequestHandler):
    hits = 0

    def do_GET(self) -> None:  # noqa: N802
        type(self).hits += 1
        self.send_response(200)
        self.end_headers()

    def log_message(self, *args: object) -> None:
        pass


def test_chromium_printing_is_offline_and_without_javascript() -> None:
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _CountingHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        url = f"http://127.0.0.1:{server.server_address[1]}"
        html = (f"<!doctype html><html><head><link rel='stylesheet' href='{url}/style.css'>"
                f"<style>@font-face{{font-family:x;src:url({url}/font.woff)}}body{{font-family:x}}</style></head>"
                f"<body><img src='{url}/image.png'><p id='p'>Texte statique</p>"
                "<script>document.getElementById('p').textContent='Texte modifié par JavaScript'</script>"
                f"<iframe src='{url}/frame'></iframe></body></html>")
        pdf, pages, _ = print_variants([html])
    finally:
        server.shutdown()
        server.server_close()
    assert pages == 1 and _CountingHandler.hits == 0
    from io import BytesIO

    text = PdfReader(BytesIO(pdf)).pages[0].extract_text()
    assert "Texte statique" in text and "JavaScript" not in text


def _descendants(pid: int) -> set[int]:
    children: dict[int, list[int]] = {}
    for entry in Path("/proc").iterdir():
        if not entry.name.isdigit():
            continue
        try:
            stat = (entry / "stat").read_text()
        except OSError:
            continue
        parent = int(stat.rsplit(")", 1)[1].split()[1])
        children.setdefault(parent, []).append(int(entry.name))
    found: set[int] = set()
    pending = [pid]
    while pending:
        for child in children.get(pending.pop(), []):
            if child not in found:
                found.add(child)
                pending.append(child)
    return found


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="lecture de /proc")
def test_browser_leaves_no_process_behind(tmp_path: Path) -> None:
    before = _descendants(os.getpid())
    _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "sobre-ats")
    with pytest.raises(RenderError):
        _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "sobre-ats", "CON")
    assert _descendants(os.getpid()) - before == set()


def test_outputs_are_written_atomically_without_leftovers(tmp_path: Path) -> None:
    existing = tmp_path / "cv.pdf"
    existing.write_bytes(b"ancien")
    result = _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "sobre-ats")
    assert result.pdf_path == existing and existing.read_bytes().startswith(b"%PDF")
    assert sorted(path.name for path in tmp_path.iterdir()) == ["cv.docx", "cv.pdf"]


@pytest.mark.parametrize("base_name", ["../cv", "CV:Entreprise", "CON", "nom.", "", "a/b", "lpt1.cv"])
def test_non_portable_output_name_is_rejected(tmp_path: Path, base_name: str) -> None:
    with pytest.raises(RenderError, match="Windows|réservé|portable"):
        _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "sobre-ats", base_name)
    assert list(tmp_path.iterdir()) == []


def test_invalid_requests_are_controlled_errors(tmp_path: Path) -> None:
    composition, identity = _load("stage_cyber"), _load("identity")
    with pytest.raises(RenderError, match="n'existe pas"):
        render_cv(composition, identity, template={"kind": "jobradar", "id": "fantaisie"}, template_map=None,
                  out_dir=tmp_path, base_name="cv")
    with pytest.raises(RenderError, match="cartographie"):
        render_cv(composition, identity, template={"kind": "original", "id": "src-1"}, template_map=None,
                  out_dir=tmp_path, base_name="cv")
    with pytest.raises(RenderError, match="composition"):
        render_cv(_load("lettre_fr"), identity, template={"kind": "jobradar", "id": "sobre-ats"},
                  template_map=None, out_dir=tmp_path, base_name="cv")
    with pytest.raises(RenderError, match="lettre"):
        render_letter(composition, identity, template={"kind": "jobradar", "id": "sobre-ats"}, template_map=None,
                      out_dir=tmp_path, base_name="lettre")
    blocker = tmp_path / "fichier"
    blocker.write_text("x")
    with pytest.raises(RenderError, match="dossier"):
        _render_cv(blocker / "sous-dossier", composition, identity, "sobre-ats")


def test_missing_browser_is_a_controlled_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jobradar.render import chromium

    def no_browser(_playwright: object) -> None:
        raise RenderError("Aucun navigateur Chromium, Chrome ou Edge utilisable n'est disponible pour produire le PDF.")

    monkeypatch.setattr(chromium, "_launch", no_browser)
    with pytest.raises(RenderError, match="navigateur"):
        _render_cv(tmp_path, _load("stage_cyber"), _load("identity"), "sobre-ats")
    assert list(tmp_path.iterdir()) == []


def test_a_trailing_blank_pdf_page_is_removed() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    raw = BytesIO()
    writer.write(raw)
    assert len(PdfReader(BytesIO(remove_trailing_blank_page(raw.getvalue()))).pages) == 1
